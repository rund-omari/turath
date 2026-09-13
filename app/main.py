import base64
import json
import os
import shutil
import sqlite3
import subprocess
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from fastapi import FastAPI, Request, UploadFile, File, Form, HTTPException
from fastapi.responses import HTMLResponse, FileResponse, RedirectResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates

BASE_DIR = Path(__file__).resolve().parent.parent
APP_DIR = BASE_DIR / "app"
DATA_DIR = Path(os.getenv("DATA_DIR", BASE_DIR / "data"))
SUBMISSIONS_DIR = DATA_DIR / "submissions"
DB_PATH = DATA_DIR / "submissions.db"
TEMPLATE_PATH = APP_DIR / "templates_docs" / "official_template.docx"
SCHEMA_PATH = APP_DIR / "form_schema.json"

DATA_DIR.mkdir(parents=True, exist_ok=True)
SUBMISSIONS_DIR.mkdir(parents=True, exist_ok=True)

app = FastAPI(title=os.getenv("APP_TITLE", "استمارة إلكترونية"))
app.mount("/static", StaticFiles(directory=APP_DIR / "static"), name="static")
templates = Jinja2Templates(directory=APP_DIR / "templates")


def load_schema() -> dict[str, Any]:
    return json.loads(SCHEMA_PATH.read_text(encoding="utf-8"))


def init_db() -> None:
    with sqlite3.connect(DB_PATH) as conn:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS submissions (
                id TEXT PRIMARY KEY,
                created_at TEXT NOT NULL,
                data_json TEXT NOT NULL,
                files_json TEXT NOT NULL,
                docx_path TEXT,
                pdf_path TEXT
            )
            """
        )
        conn.commit()


def basic_auth_ok(request: Request) -> bool:
    expected_user = os.getenv("ADMIN_USER", "admin")
    expected_password = os.getenv("ADMIN_PASSWORD", "change-this-password")
    auth = request.headers.get("Authorization", "")
    if not auth.startswith("Basic "):
        return False
    try:
        decoded = base64.b64decode(auth.split(" ", 1)[1]).decode("utf-8")
        user, password = decoded.split(":", 1)
    except Exception:
        return False
    return user == expected_user and password == expected_password


def require_admin(request: Request) -> None:
    if not basic_auth_ok(request):
        raise HTTPException(
            status_code=401,
            detail="Unauthorized",
            headers={"WWW-Authenticate": 'Basic realm="Admin"'},
        )


def flatten_fields(schema: dict[str, Any]) -> list[dict[str, Any]]:
    out = []
    for section in schema.get("sections", []):
        out.extend(section.get("fields", []))
    return out


def make_reference() -> str:
    return f"MOC-{datetime.now().strftime('%Y%m%d')}-{uuid.uuid4().hex[:6].upper()}"


def create_fallback_docx(path: Path, schema: dict[str, Any], values: dict[str, Any], reference: str) -> None:
    doc = Document()
    p = doc.add_paragraph()
    p.alignment = 2
    r = p.add_run(schema.get("title", "الاستمارة"))
    r.bold = True
    p2 = doc.add_paragraph()
    p2.alignment = 2
    p2.add_run(f"رقم الطلب: {reference}")
    for section in schema.get("sections", []):
        h = doc.add_paragraph()
        h.alignment = 2
        rr = h.add_run(section.get("title", ""))
        rr.bold = True
        for field in section.get("fields", []):
            if field.get("type") == "file":
                continue
            p = doc.add_paragraph()
            p.alignment = 2
            value = values.get(field["name"], "")
            p.add_run(f"{field.get('label', field['name'])}: {value}")
    doc.save(path)


def _replace_in_paragraph(paragraph, context: dict[str, Any]) -> None:
    # Rebuild the paragraph text only when a placeholder is present.
    # This keeps the implementation dependency-light and works well for official templates
    # where placeholders are typed as plain text like {{ applicant_name }}.
    full_text = "".join(run.text for run in paragraph.runs)
    replaced = full_text
    for key, value in context.items():
        replaced = replaced.replace("{{ " + key + " }}", str(value))
        replaced = replaced.replace("{{" + key + "}}", str(value))
    if replaced != full_text:
        if paragraph.runs:
            paragraph.runs[0].text = replaced
            for run in paragraph.runs[1:]:
                run.text = ""
        else:
            paragraph.add_run(replaced)


def _replace_in_table(table, context: dict[str, Any]) -> None:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                _replace_in_paragraph(paragraph, context)
            for nested in cell.tables:
                _replace_in_table(nested, context)


def _set_cell(cell, text: str) -> None:
    # Preserve the official table structure while putting the answer in its designated answer cell.
    cell.text = str(text or "")
    for p in cell.paragraphs:
        p.alignment = 2


def render_docx(template_path: Path, output_path: Path, context: dict[str, Any], schema: dict[str, Any], reference: str) -> None:
    if not template_path.exists():
        create_fallback_docx(output_path, schema, context, reference)
        return
    doc = Document(template_path)
    t = doc.tables[0]
    row_map = {
        1:"researcher_name", 2:"card_number", 3:"collection_date_place", 4:"governorate",
        6:"element_name", 7:"other_names", 9:"geographic_location", 10:"practicing_communities",
        11:"element_description", 12:"social_functions", 13:"practice_context", 15:"current_status",
        16:"formal_safeguarding", 17:"informal_safeguarding", 18:"proposed_safeguarding"
    }
    for row, key in row_map.items():
        _set_cell(t.cell(row,0), context.get(key,""))
    selected = set(x.strip() for x in context.get("classification","").split("،") if x.strip())
    class_opts = ['التقاليد وأشكال التعبير الشفهي، بما في ذلك اللغة.','فنون وتقاليد أداء العروض','الممارسات الاجتماعية والطقوس والاحتفالات.','المعارف والممارسات المتعلقة بالطبيعة والكون.','المهارات المرتبطة بالفنون الحرفية التقليدية.']
    _set_cell(t.cell(8,0), "\n".join(("☒ " if x in selected else "☐ ")+x for x in class_opts))
    _set_cell(t.cell(14,0), "1. طرق رسمية (الحكومة)\n"+context.get("transmission_formal","")+"\n\n2. طرق غير رسمية:\n"+context.get("transmission_informal",""))
    sdgs=['الأمن الغذائي','الرعاية الصحية','التعليم الجيد','المساواة بين الجنسين','التنمية الاقتصادية الشاملة','الاستدامة البيئية بما في ذلك تغير المناخ','السلام والتماسك الاجتماعي']
    sel=set(x.strip() for x in context.get("sdg_relation","").split("،") if x.strip())
    _set_cell(t.cell(19,0), "أمثلة:\n"+"\n".join(("☒ " if x in sel else "☐ ")+x for x in sdgs))
    # Practitioners table nested in row 20, first cell: header + 4 data rows.
    nested = t.cell(20,0).tables[0] if t.cell(20,0).tables else None
    if nested:
        for i in range(1, min(5,len(nested.rows))):
            vals=[context.get(f"practitioner_{i}_organization",""),context.get(f"practitioner_{i}_person",""),context.get(f"practitioner_{i}_phone",""),context.get(f"practitioner_{i}_address",""),context.get(f"practitioner_{i}_email",""),context.get(f"practitioner_{i}_social","")]
            # Visual order in the source table is RTL; underlying cells run left-to-right: social,email,address,phone,person,organization.
            for j,val in enumerate(reversed(vals)):
                _set_cell(nested.cell(i,j),val)
    narrator = "\n".join([f"الاسم: {context.get('narrator_name','')}",f"الصفة: {context.get('narrator_role','')}",f"الجنس: {context.get('narrator_gender','')}",f"رقم الهاتف: {context.get('narrator_phone','')}",f"العنوان: {context.get('narrator_address','')}",f"البريد الإلكتروني: {context.get('narrator_email','')}"])
    _set_cell(t.cell(21,0), narrator)
    doc.save(output_path)


def convert_to_pdf(docx_path: Path, out_dir: Path) -> Path | None:
    try:
        subprocess.run(
            ["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(docx_path)],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            timeout=90,
        )
        pdf = out_dir / f"{docx_path.stem}.pdf"
        return pdf if pdf.exists() else None
    except Exception:
        return None


init_db()


@app.get("/health")
def health():
    return {"ok": True}


@app.get("/", response_class=HTMLResponse)
def form_page(request: Request):
    schema = load_schema()
    return templates.TemplateResponse("form.html", {"request": request, "schema": schema})


@app.post("/submit", response_class=HTMLResponse)
async def submit(request: Request):
    schema = load_schema()
    fields = flatten_fields(schema)
    form = await request.form()
    values: dict[str, Any] = {}

    missing = []
    for field in fields:
        name = field["name"]
        if field.get("type") == "file":
            continue
        if field.get("type") == "checkbox":
            value = "، ".join(str(x).strip() for x in form.getlist(name) if str(x).strip())
        else:
            value = str(form.get(name, field.get("default", ""))).strip()
        values[name] = value
        if field.get("required") and not value:
            missing.append(field.get("label", name))

    if missing:
        return templates.TemplateResponse(
            "error.html",
            {"request": request, "message": "يرجى تعبئة الحقول المطلوبة: " + "، ".join(missing)},
            status_code=422,
        )

    reference = make_reference()
    submission_dir = SUBMISSIONS_DIR / reference
    uploads_dir = submission_dir / "uploads"
    submission_dir.mkdir(parents=True, exist_ok=True)
    uploads_dir.mkdir(parents=True, exist_ok=True)

    saved_files = []
    for field in fields:
        if field.get("type") != "file":
            continue
        items = form.getlist(field["name"])
        for item in items:
            if not hasattr(item, "filename") or not item.filename:
                continue
            safe_name = Path(item.filename).name.replace("..", "_")
            target = uploads_dir / safe_name
            content = await item.read()
            target.write_bytes(content)
            saved_files.append(str(target.relative_to(DATA_DIR)))

    values["reference_number"] = reference
    values["submission_date"] = datetime.now().astimezone().strftime("%Y-%m-%d %H:%M")

    docx_path = submission_dir / f"{reference}.docx"
    render_docx(TEMPLATE_PATH, docx_path, values, schema, reference)
    pdf_path = convert_to_pdf(docx_path, submission_dir)

    created_at = datetime.now(timezone.utc).isoformat()
    with sqlite3.connect(DB_PATH) as conn:
        conn.execute(
            "INSERT INTO submissions (id, created_at, data_json, files_json, docx_path, pdf_path) VALUES (?, ?, ?, ?, ?, ?)",
            (
                reference,
                created_at,
                json.dumps(values, ensure_ascii=False),
                json.dumps(saved_files, ensure_ascii=False),
                str(docx_path.relative_to(DATA_DIR)),
                str(pdf_path.relative_to(DATA_DIR)) if pdf_path else None,
            ),
        )
        conn.commit()

    return templates.TemplateResponse(
        "success.html",
        {"request": request, "reference": reference},
    )


@app.get("/admin", response_class=HTMLResponse)
def admin(request: Request):
    require_admin(request)
    with sqlite3.connect(DB_PATH) as conn:
        conn.row_factory = sqlite3.Row
        rows = conn.execute("SELECT * FROM submissions ORDER BY created_at DESC").fetchall()
    items = []
    for row in rows:
        item = dict(row)
        item["data"] = json.loads(item.pop("data_json"))
        item["files"] = json.loads(item.pop("files_json"))
        items.append(item)
    return templates.TemplateResponse("admin.html", {"request": request, "items": items})


@app.get("/admin/files/{reference}/{kind}")
def admin_file(reference: str, kind: str, request: Request):
    require_admin(request)
    if kind not in {"docx", "pdf"}:
        raise HTTPException(404)
    with sqlite3.connect(DB_PATH) as conn:
        row = conn.execute("SELECT docx_path, pdf_path FROM submissions WHERE id = ?", (reference,)).fetchone()
    if not row:
        raise HTTPException(404)
    rel = row[0] if kind == "docx" else row[1]
    if not rel:
        raise HTTPException(404)
    path = DATA_DIR / rel
    if not path.exists():
        raise HTTPException(404)
    return FileResponse(path, filename=path.name)
